# -*- coding: utf-8 -*-
"""
ai_kb.ingest — 知识库多格式入库解析（TASK-010 包B）
职责：后缀白名单校验 → 按格式解析为纯文本 → 清洗（剥HTML标签/控制字符）→ 返回摘要。
不做切片、不碰索引（那是 pipeline/retriever 的事）；不落库。
数值规格（任务书写死）：单文件 20MB；白名单 .txt .md .docx .pdf；
扫描版检测 = 平均每页提取字符 < 30 → scanned_pdf_unsupported。
"""

import io
import os
import re

MAX_FILE_SIZE = 20 * 1024 * 1024          # 20MB
ALLOWED_EXT = ('.txt', '.md', '.docx', '.pdf')
SCANNED_PAGE_CHAR_THRESHOLD = 30          # 平均每页 <30 字符判扫描版


class IngestError(Exception):
    """code: unsupported_type / too_large / encoding_error / parse_error / scanned_pdf_unsupported"""

    def __init__(self, code, msg):
        super().__init__(msg)
        self.code = code
        self.msg = msg


def _ext_of(filename):
    return os.path.splitext(filename or '')[1].lower()


def clean_text(text):
    """剥 HTML 标签（红线3：入库文本先剥标签）+ 去控制字符（保留换行制表）+ 压缩空行。"""
    text = re.sub(r'<[^>]{0,200}>', '', text)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def _decode_bytes(raw):
    for enc in ('utf-8-sig', 'utf-8', 'gbk'):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    raise IngestError('encoding_error', '无法识别文本编码（支持 utf-8/gbk）')


def _parse_txt(raw, filename):
    ext = _ext_of(filename)
    if ext in ('.md', '.txt'):
        return _decode_bytes(raw), None   # 页数=None
    raise IngestError('unsupported_type', '不支持的文件类型: %s' % ext)


def _parse_docx(raw):
    try:
        import docx  # python-docx
        d = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs]
        for tb in d.tables:                # 表格转管道行，交给切片器当表格单元
            for row in tb.rows:
                parts.append('| ' + ' | '.join(c.text.strip() for c in row.cells) + ' |')
        return '\n'.join(parts)
    except IngestError:
        raise
    except Exception as e:
        raise IngestError('parse_error', 'docx 解析失败（加密或损坏）: %s' % type(e).__name__)


def _parse_pdf(raw):
    try:
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(raw))
        pages = []
        for pg in r.pages:
            try:
                pages.append(pg.extract_text() or '')
            except Exception:
                pages.append('')
        n = len(pages) or 1
        total = sum(len(p.strip()) for p in pages)
        if total / n < SCANNED_PAGE_CHAR_THRESHOLD:
            raise IngestError('scanned_pdf_unsupported',
                              '暂不支持扫描件PDF（平均每页不足%d字），请先转成文字版' % SCANNED_PAGE_CHAR_THRESHOLD)
        return '\n\n'.join(pages)          # 页间插换行
    except IngestError:
        raise
    except Exception as e:
        raise IngestError('parse_error', 'PDF 解析失败: %s' % type(e).__name__)


def parse_upload(filename, raw_bytes):
    """入口：返回 {text, pages, chars}。任何失败抛 IngestError(code)。"""
    if not raw_bytes:
        raise IngestError('parse_error', '空文件')
    if len(raw_bytes) > MAX_FILE_SIZE:
        raise IngestError('too_large', '文件超过20MB上限')
    ext = _ext_of(filename)
    if ext not in ALLOWED_EXT:
        raise IngestError('unsupported_type', '只支持 .txt .md .docx .pdf，收到: %s' % (ext or '无后缀'))
    if ext in ('.txt', '.md'):
        text, pages = _parse_txt(raw_bytes, filename)
    elif ext == '.docx':
        text, pages = _parse_docx(raw_bytes), None
    else:
        text, pages = _parse_pdf(raw_bytes), None
    text = clean_text(text)
    if len(text) < 50:
        raise IngestError('parse_error', '解析后正文不足50字，疑似空文档或乱码')
    return {'text': text, 'pages': pages, 'chars': len(text)}
